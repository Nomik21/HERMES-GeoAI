from pathlib import Path
import csv
from openpyxl import load_workbook
from collections import Counter

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATASET_ROOT = PROJECT_ROOT / "data" / "hermes_benchmark"

MANIFEST_CSV = PROJECT_ROOT / "evaluation" / "hermes_dataset_manifest.csv"
OUTPUT_CSV = PROJECT_ROOT / "evaluation" / "hermes_object_groundtruth.csv"

OBJECT_COLUMNS = [
    "Obj_Debris",
    "Obj_FallenTree",
    "Obj_Flooded",
    "Obj_DamagedBldg",
    "Obj_DownedLines",
]


def normalize_cell(value):
    if value is None:
        return ""

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)

    if isinstance(value, int):
        return str(value)

    text = str(value).strip()

    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]

    return text


def basename(path_text):
    return str(path_text).replace("\\", "/").split("/")[-1]


def read_sheet_rows(xlsx_path: Path):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active

    iterator = ws.iter_rows(values_only=True)
    header = next(iterator)
    header = [str(h) for h in header]

    for raw_row in iterator:
        if raw_row is None:
            continue

        row = dict(zip(header, raw_row))

        if all(v is None for v in row.values()):
            continue

        yield row


# ---------------------------------------------------------------------
# Build manifest lookup by scenario + image filenames
# ---------------------------------------------------------------------
with MANIFEST_CSV.open("r", encoding="utf-8-sig", newline="") as f:
    manifest_rows = list(csv.DictReader(f))

manifest_by_images = {}

for row in manifest_rows:
    scenario = row["scenario"]

    if scenario not in {"cross_view_svi_rsi", "pre_post_svi"}:
        continue

    image_1 = basename(row["image_path_1"])
    image_2 = basename(row["image_path_2"])

    key = (scenario, image_1, image_2)
    manifest_by_images[key] = row["case_id"]


rows = []
unmatched = []


def add_object_row(
    *,
    scenario,
    dataset,
    raw_case_id,
    input_1,
    input_2,
    source_row,
):
    key = (scenario, basename(input_1), basename(input_2))
    canonical_case_id = manifest_by_images.get(key)

    if canonical_case_id is None:
        unmatched.append(
            {
                "scenario": scenario,
                "raw_case_id": raw_case_id,
                "input_1": input_1,
                "input_2": input_2,
                "lookup_key": key,
            }
        )
        canonical_case_id = raw_case_id

    out = {
        "case_id": canonical_case_id,
        "object_gt_raw_id": raw_case_id,
        "scenario": scenario,
        "dataset": dataset,
        "input_1": basename(input_1),
        "input_2": basename(input_2),
    }

    for col in OBJECT_COLUMNS:
        out[col] = int(float(source_row.get(col, 0) or 0))

    rows.append(out)


# ---------------------------------------------------------------------
# CVIAN / cross-view object ground truth
# ---------------------------------------------------------------------
cvian_xlsx = DATASET_ROOT / "cvian-object-groundtruth.xlsx"

for row in read_sheet_rows(cvian_xlsx):
    raw_case_id = normalize_cell(row["Location_ID"])
    input_1 = normalize_cell(row["Input_RSI"])
    input_2 = normalize_cell(row["Input_SVI"])

    add_object_row(
        scenario="cross_view_svi_rsi",
        dataset="CVIAN_selected_50",
        raw_case_id=raw_case_id,
        input_1=input_1,
        input_2=input_2,
        source_row=row,
    )


# ---------------------------------------------------------------------
# Pre/post SVI object ground truth
# ---------------------------------------------------------------------
prepost_xlsx = DATASET_ROOT / "pre&post-object-groundtruth.xlsx"

for row in read_sheet_rows(prepost_xlsx):
    raw_case_id = normalize_cell(row["Pair_ID"])
    input_1 = normalize_cell(row["Pre_Image_File"])
    input_2 = normalize_cell(row["Post_Image_File"])

    add_object_row(
        scenario="pre_post_svi",
        dataset="pre_post_SVI_50_50_50",
        raw_case_id=raw_case_id,
        input_1=input_1,
        input_2=input_2,
        source_row=row,
    )


# ---------------------------------------------------------------------
# Write unified object ground truth CSV
# ---------------------------------------------------------------------
OUTPUT_CSV.parent.mkdir(parents=True, exist_ok=True)

fieldnames = [
    "case_id",
    "object_gt_raw_id",
    "scenario",
    "dataset",
    "input_1",
    "input_2",
] + OBJECT_COLUMNS

with OUTPUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames)
    writer.writeheader()
    writer.writerows(rows)

print(f"Object ground truth written to: {OUTPUT_CSV}")
print(f"Total rows: {len(rows)}")
print(f"Unmatched rows: {len(unmatched)}")

if unmatched:
    print("\nFirst unmatched rows:")
    for item in unmatched[:20]:
        print(item)

print("\nRows by scenario:")
for scenario, count in Counter(r["scenario"] for r in rows).items():
    print(f"  {scenario}: {count}")

print("\nPositive labels by object:")
for col in OBJECT_COLUMNS:
    count = sum(int(r[col]) for r in rows)
    print(f"  {col}: {count}")
