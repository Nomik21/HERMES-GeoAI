from pathlib import Path
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATASET_ROOT = PROJECT_ROOT / "data" / "hermes_benchmark"

xlsx_files = [
    DATASET_ROOT / "cvian-object-groundtruth.xlsx",
    DATASET_ROOT / "pre&post-object-groundtruth.xlsx",
]

for xlsx_path in xlsx_files:
    print("\n" + "=" * 100)
    print(f"FILE: {xlsx_path}")
    print("=" * 100)

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)

    for ws in wb.worksheets:
        print(f"\nSHEET: {ws.title}")
        print(f"Rows: {ws.max_row}, Columns: {ws.max_column}")

        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)

        print("\nHEADER:")
        print(header)

        print("\nFIRST 5 DATA ROWS:")
        for i, row in enumerate(rows, start=1):
            print(row)
            if i >= 5:
                break
